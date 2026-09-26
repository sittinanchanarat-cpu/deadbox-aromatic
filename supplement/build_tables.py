#!/usr/bin/env python
"""Supplementary Tables S1-S6 as one Excel workbook (NAR Supplementary Data).

Every sheet is built from the stage results already in the repository; nothing
is recomputed except the census summaries of Table S3, which repeat the counting
of figures/fig1_definition.py (panel B) and figures/fig3_aromatic_anchor.py
(panel A). Tables are numbered by first citation in Materials and Methods.

A README sheet lists the tables and defines every column.

Output
  supplement/Supplementary_Tables_S1-S6.xlsx
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "supplement" / "Supplementary_Tables_S1-S6.xlsx"
S0, S1, S2, S3, S4 = (ROOT / f"stage{i}" / "results" for i in range(5))
FONT = "Arial"


def tsv(path):
    return pd.read_csv(path, sep="\t", low_memory=False)


def pick(df, cols):
    """Select and rename columns; cols = [(source, header, definition), ...]."""
    d = df[[c for c, _, _ in cols]].copy()
    d.columns = [h for _, h, _ in cols]
    return d, [(h, text) for _, h, text in cols]


# --------------------------------------------------------------------- S1 structures
def s1a():
    d = tsv(S0 / "sites_labelled.tsv")
    return pick(d, [
        ("site", "Site", "Nucleotide site: PDB entry : chain of the nucleotide : ligand code and residue number"),
        ("pdb", "PDB", "PDB entry"),
        ("chain", "Helicase chain", "Chain of the helicase that binds the nucleotide"),
        ("ligand", "Nucleotide", "Chemical Component Dictionary code of the adenine nucleotide"),
        ("key", "UniProt", "UniProt accession of the helicase (entity description when no accession is given)"),
        ("protein", "Protein", "Protein name (UniProt)"),
        ("family", "Family", "SF2 helicase family (hand-curated; Table S1B)"),
        ("stacker", "Adenine stacker", "RecA1 aromatic stacked on the adenine (ring centroids <= 5.0 Å, interplanar angle <= 30°)"),
        ("stacker_dist", "Stacker-adenine distance (Å)", "Ring centroid to purine centroid"),
        ("stacker_angle", "Stacker-adenine angle (°)", "Angle between the aromatic and purine planes"),
        ("q_gln", "Adenine-reading Gln", "Glutamine whose side chain hydrogen bonds adenine N6/N7 (chain:residue)"),
        ("q_contacts", "Gln-adenine H-bonds (Å)", "Gln OE1/NE2 distances to adenine N6 and N7"),
        ("walker_lys", "Motif I Lys", "Walker A lysine whose NZ contacts the phosphates (chain:residue)"),
        ("stacker_to_q", "Stacker to Gln (residues)", "Sequence separation, entity numbering"),
        ("stacker_to_walkerK", "Stacker to motif I Lys (residues)", "Sequence separation, entity numbering"),
        ("q_to_walkerK", "Gln to motif I Lys (residues)", "Sequence separation, entity numbering"),
        ("ua_state", "Anchor state", "inserted: an aromatic 38-60 residues before the motif I Lys has its ring centroid <= 9.5 Å from "
                                     "the purine centroid and an atom <= 5.0 Å from the nucleotide; displaced: aromatics in that "
                                     "window but none meets both; absent: no aromatic in a mostly modelled window; not_modelled: "
                                     "half or more of the window unmodelled; no_walkerK: no motif I Lys at the site"),
        ("ua", "Anchor residue", "Inserted aromatic (or, if none, the window aromatic closest to the purine); author numbering"),
        ("ua_aa", "Anchor amino acid", "One-letter code of the anchor residue"),
        ("ua_to_walkerK", "Anchor to motif I Lys (residues)", "Sequence separation, entity numbering"),
        ("anchor", "Register landmark", "Landmark used to place the window: Q (Q-motif Gln, taken as motif I Lys - 23) or walkerK"),
        ("ua_ring_to_adenine", "Anchor ring-purine distance (Å)", "Ring centroid to purine centroid"),
        ("ua_min_to_nucleotide", "Anchor-nucleotide closest distance (Å)", "Closest atom pair"),
        ("ua_bfactor_z", "Anchor B-factor (z)", "Mean anchor B-factor as a z-score over the protein chain"),
        ("loop_missing", "Unmodelled window residues", "Residues of the 23-residue window absent from the model"),
        ("n_window_aromatics", "Window aromatics", "F, W, Y or H residues in the window"),
        ("window_seq", "Window sequence", "Motif I Lys - 60 to - 38, entity sequence; '-' = unmodelled"),
    ])


def s1b():
    d = tsv(ROOT / "stage0" / "data" / "family_curation.tsv")
    d["family"] = d.family_curated.fillna(d.family_auto)
    d["checked"] = np.where(d.family_curated.notna(), "corrected by hand", "keyword label confirmed")
    return pick(d, [
        ("key", "UniProt", "UniProt accession (entity description when no accession is given)"),
        ("uniprot_name", "UniProt name", "Recommended name"),
        ("description", "PDB entity description", "As deposited"),
        ("organism", "Organism", "Source organism"),
        ("n_entries", "PDB entries", "Nucleotide-bound entries of this protein"),
        ("example_pdbs", "Example entries", "Up to three entries"),
        ("family_auto", "Keyword family", "Family from UniProt name and Pfam keywords; ? = no keyword matched"),
        ("family", "Family (final)", "Family used in the analysis"),
        ("checked", "Curation", "Whether the keyword label was replaced by hand"),
        ("note", "Note", "Curation note"),
    ])


def s1c():
    d = tsv(S0 / "n6_geometry.tsv")
    return pick(d, [
        ("site", "Site", "As in Table S1A"),
        ("pdb", "PDB", "PDB entry"),
        ("protein", "Protein", "Protein name"),
        ("ua", "Anchor", "Inserted anchor residue, author numbering"),
        ("nucleotide", "Nucleotide", "Ligand code"),
        ("state", "Nucleotide state", "ATP-like (non-hydrolysable analogues), ADP (+TS-mimic when a transition-state mimic is present), AMP"),
        ("rna_present", "RNA in entry", "Whether the entry contains RNA (proxy for the closed state)"),
        ("n6_to_ring", "N6-ring centroid (Å)", "Adenine N6 to anchor ring centroid"),
        ("approach_angle", "Approach angle (°)", "Ring normal vs centroid-to-N6 vector; ~0° face-on, ~90° edge-on"),
        ("n6_to_ring_atom", "Closest ring atom", "Anchor ring atom closest to N6 and its distance (Å)"),
        ("min_ring_atom_dist", "N6-closest ring atom (Å)", "Distance to that atom"),
        ("q_hbond_n6", "Gln-N6 (Å)", "Q-motif Gln side-chain O/N to adenine N6"),
    ])


# --------------------------------------------------------------------- S2 partner complexes
PARTNER_COLS = [
    ("pdb", "PDB", "PDB entry"),
    ("chain", "Chain", "DEAD-box chain"),
    ("protein", "Protein", "Entity description"),
    ("uniprot", "UniProt", "UniProt accession"),
    ("organism", "Organism", "Source organism"),
    ("method", "Method", "Experimental method"),
    ("resolution", "Resolution (Å)", "As deposited"),
    ("ua_aa", "Anchor amino acid", "Aromatic 21-30 residues before the Q-motif Gln, nearest to Q-25"),
    ("ua_seq", "Anchor position (entity)", "Anchor position in the entity sequence"),
    ("q_seq", "Q position (entity)", "Q-motif Gln position in the entity sequence"),
    ("ua_auth", "Anchor residue", "Anchor, author numbering"),
    ("loop_modelled", "Loop residues modelled", "Of the 12-residue anchor loop UA-8 to UA+3"),
    ("ua_sidechain", "Anchor side chain modelled", "Whether the ring atoms are present"),
    ("ua_to_qaromatic", "Anchor-stacker ring distance (Å)", "Anchor ring centroid to Q-motif stacker ring centroid"),
    ("ua_ca_to_q_ca", "Anchor Cα-Gln Cα (Å)", "Used when side chains are not built"),
    ("ua_state", "Anchor state", "inserted: ring centroid <= 8.0 Å (8.5 Å for Trp) from the Q-motif stacker ring; displaced: further; "
                                 "(bb): side chain not built, called from the anchor Cα-Gln Cα distance (inserted <= 10.0 Å, "
                                 "10.6 Å for Trp; displaced >= 12.0 Å; ambiguous between); not_modelled: anchor absent from the "
                                 "model; no_UA: no aromatic at Q-21 to Q-30"),
    ("ua_ring_to_N6", "Anchor ring-N6 (Å)", "Ring centroid to adenine N6 when a nucleotide is bound"),
    ("ua_partner", "Partners at anchor", "Other entities (cognate copy only) with an atom <= 4.5 Å from the anchor"),
    ("ua_self", "Self copies at anchor", "Other copies of the same entity within 4.5 Å of the anchor"),
    ("loop_partner", "Partners at loop", "Other entities within 4.5 Å of the anchor loop (UA-8 to UA+3)"),
    ("loop_self", "Self copies at loop", "Other copies of the same entity within 4.5 Å of the loop"),
    ("n_ua_partner_atoms", "Partner atoms at anchor", "Partner atoms within 4.5 Å of the anchor"),
    ("n_loop_partner_atoms", "Partner atoms at loop", "Partner atoms within 4.5 Å of the loop"),
]


def s2a():
    return pick(tsv(S0 / "partner_contacts.tsv"), PARTNER_COLS)


def s2b():
    d = tsv(S0 / "loop_summary.tsv")
    return pick(d, [
        ("protein", "Protein", "Panel protein"),
        ("pdb", "PDB", "PDB entry"),
        ("chain", "Chain", "Helicase chain"),
        ("label", "Partner state", "free, or the bound partner"),
        ("resolution", "Resolution (Å)", "As deposited"),
        ("ref", "Reference chain", "Free chain whose numbering is used; every chain is superposed on it"),
        ("core_n", "Core Cα used", "RecA1 core Cα atoms (UA+8 to UA+190) kept after outlier rejection (> 2 Å)"),
        ("core_rmsd", "Core RMSD (Å)", "Superposition RMSD on those atoms"),
        ("window_modelled", "Window residues modelled", "Of the 19 residues UA-12 to UA+6"),
        ("ua", "Anchor", "Anchor residue, reference numbering"),
        ("ua_ring_shift_vs_ref", "Anchor ring shift vs reference (Å)", "Ring centroid displacement after core superposition"),
        ("ua_chi1", "Anchor χ1 (°)", "Side-chain torsion"),
        ("ua_chi2", "Anchor χ2 (°)", "Side-chain torsion"),
        ("nearest_free", "Nearest free chain", "Free chain of the same protein from another entry with the lowest loop RMSD"),
        ("loop_rmsd_nearest_free", "Loop Cα RMSD to nearest free (Å)", "Window UA-12 to UA+6"),
        ("loop_n_compared", "Loop Cα compared", "Window residues modelled in both chains"),
        ("ua_ring_shift_nearest_free", "Anchor ring shift vs nearest free (Å)", "Ring centroid displacement"),
        ("window_touched", "Window residues contacted", "Window residues with a partner atom <= 4.5 Å (cognate copy, no symmetry mates)"),
        ("window_buried_A2", "Window area buried (Å²)", "SASA of the window lost on partner binding (Shrake-Rupley)"),
        ("loop_partners", "Loop partners", "Partner entities contacting the window"),
    ])


def s2c():
    d = tsv(S0 / "partner_contacts.tsv")
    d = d[d.ua_state.str.startswith(("displaced", "ambiguous"))]
    return pick(d, [c for c in PARTNER_COLS if c[0] in (
        "pdb", "chain", "protein", "uniprot", "organism", "method", "resolution", "ua_auth", "loop_modelled",
        "ua_sidechain", "ua_to_qaromatic", "ua_ca_to_q_ca", "ua_state", "ua_partner", "loop_partner")])


# --------------------------------------------------------------------- S3 sequence census
def census():
    c = tsv(S1 / "census_balanced_assigned.tsv")
    c = c[c.q_ok.fillna(False).astype(bool)].copy()
    c["anchor"] = c.ua_aa.where(c.ua_found.fillna(False).astype(bool), "-").where(lambda s: s.isin(list("FWY-")), "-")
    e = c[(c.domain == "Eukaryota") & ~c.subfamily_final.isin(["ambiguous", "no_hit", "unassigned"])]
    e = e.sort_values("assign_how", key=lambda s: s.ne("label")).drop_duplicates(["genus", "subfamily_final"])
    return c, e


def s3a(c, e):
    t = pd.crosstab(e.subfamily_final, e.anchor).reindex(columns=list("FWY-"), fill_value=0)
    t.loc["Bacteria (all sequences)"] = c[c.domain == "Bacteria"].anchor.value_counts().reindex(list("FWY-")).fillna(0)
    n = t.sum(axis=1)
    d = pd.DataFrame({"group": t.index, "n": n.astype(int).values})
    for aa, lab in (("F", "Phe"), ("W", "Trp"), ("Y", "Tyr"), ("-", "none")):
        d[lab] = t[aa].astype(int).values
    for aa, lab in (("F", "Phe"), ("W", "Trp"), ("Y", "Tyr")):
        d[f"{lab} %"] = (100 * t[aa] / n).round(1).values
    d["arom %"] = (100 * t[["F", "W", "Y"]].sum(axis=1) / n).round(1).values
    d["fig"] = np.where((d.n >= 100) | (d.group == "Bacteria (all sequences)"), "yes", "no")
    # median anchor spacing from the Q-motif Gln
    sp = e[e.anchor != "-"].groupby("subfamily_final").ua_spacing.median()
    d["spacing"] = d.group.map(sp)
    d["bac"] = d.group.str.startswith("Bacteria")
    d = d.sort_values(["bac", "fig", "Trp %", "group"], ascending=[False, False, False, True])
    return pick(d, [
        ("group", "Subfamily", "UniProt subfamily (unlabelled sequences assigned by nearest labelled homologue)"),
        ("n", "Sequences", "Eukaryotes: one sequence per genus and subfamily, so also the number of genera; Bacteria: all sequences with a Q-motif Gln"),
        ("Phe", "Phe", "Sequences with a Phe anchor"),
        ("Trp", "Trp", "Sequences with a Trp anchor"),
        ("Tyr", "Tyr", "Sequences with a Tyr anchor"),
        ("none", "No aromatic", "No aromatic in the anchor window"),
        ("Phe %", "Phe (%)", "Percentage of sequences"),
        ("Trp %", "Trp (%)", "Percentage of sequences"),
        ("Tyr %", "Tyr (%)", "Percentage of sequences"),
        ("arom %", "Aromatic (%)", "Phe + Trp + Tyr"),
        ("spacing", "Median spacing (Q-n)", "Median number of residues from the anchor to the Q-motif Gln"),
        ("fig", "In Figure 3A", "Subfamilies with at least 100 sequences, and Bacteria"),
    ])


def s3b(e):
    e = e[e.anchor != "-"]
    t = pd.crosstab(e.ua_spacing.astype(int), e.anchor).reindex(columns=list("FWY"), fill_value=0)
    d = t.reset_index().rename(columns={"ua_spacing": "spacing"})
    d["total"] = t.sum(axis=1).values
    d["pct"] = (100 * d.total / d.total.sum()).round(2)
    top = e.groupby(e.ua_spacing.astype(int)).subfamily_final.agg(
        lambda s: "; ".join(f"{k} ({v})" for k, v in s.value_counts().head(4).items()))
    d["subfamilies"] = d.spacing.map(top)
    return pick(d, [
        ("spacing", "Anchor position (Q-n)", "Residues from the anchor to the Q-motif Gln"),
        ("F", "Phe", "Sequences"),
        ("W", "Trp", "Sequences"),
        ("Y", "Tyr", "Sequences"),
        ("total", "Total", "Sequences"),
        ("pct", "Total (%)", "Percentage of eukaryotic sequences with an aromatic anchor"),
        ("subfamilies", "Main subfamilies", "Up to four subfamilies with most sequences at this spacing (count)"),
    ])


def s3c():
    d = tsv(S3 / "supergroups.tsv")
    return pick(d, [
        ("subfamily", "Subfamily", "Subfamily"),
        ("supergroup", "Supergroup", "Eukaryotic supergroup of the genus"),
        ("n", "Genera", "Genera (one sequence each)"),
        ("pct_W", "Trp (%)", "Genera with a Trp anchor"),
        ("pct_F", "Phe (%)", "Genera with a Phe anchor"),
    ])


# --------------------------------------------------------------------- S4 AlphaFold
def s4a():
    d = tsv(S2 / "af_anchor.tsv")
    return pick(d, [
        ("acc", "UniProt", "Accession; AlphaFold DB model AF-<accession>-F1 (version 6)"),
        ("subfamily", "Subfamily", "Subfamily (census assignment)"),
        ("label", "Reference", "Named yeast or human reference protein"),
        ("organism", "Organism", "Source organism"),
        ("class", "Class", "Taxonomic class"),
        ("census_anchor", "Sequence anchor", "Anchor amino acid called from sequence (Table S3)"),
        ("census_pos", "Sequence anchor position", "UniProt numbering"),
        ("status", "Status", "ok: analysed; no_model: no AlphaFold DB model; motif_not_found: RecA1 motifs for superposition not found"),
        ("motif_rmsd", "Motif RMSD (Å)", "Superposition of 19 RecA1 motif Cα atoms on Vasa (2DB3 chain A)"),
        ("qarom", "Q-motif stacker", "Aromatic at Q-7/Q-6"),
        ("qarom_to_adenine", "Stacker-adenine (Å)", "Stacker ring centroid to transplanted adenine"),
        ("struct_anchor", "Structural anchor", "Aromatic at Q-30 to Q-20 inserted at the adenine (N6-closest ring atom <= 5.0 Å and "
                                                "N6-ring centroid <= 7.0 Å) with the shortest N6 distance"),
        ("struct_aa", "Structural anchor amino acid", "'-' = no inserted aromatic"),
        ("struct_offset", "Structural anchor position (Q-n)", "Residues before the Q-motif Gln"),
        ("matches_census", "Matches sequence call", "Structural and sequence anchors are the same residue"),
        ("n6_centroid", "N6-ring centroid (Å)", "Adenine N6 to anchor ring centroid"),
        ("n6_closest_atom", "Closest ring atom", "Anchor ring atom closest to N6"),
        ("n6_closest", "N6-closest ring atom (Å)", "Distance to that atom"),
        ("approach_angle", "Approach angle (°)", "Ring normal vs centroid-to-N6 vector"),
        ("to_qarom", "Anchor-stacker ring distance (Å)", "Anchor ring centroid to Q-motif stacker ring centroid"),
        ("rsasa", "Relative SASA", "Relative side-chain solvent accessibility of the anchor in the model"),
        ("plddt", "pLDDT", "Model confidence at the anchor"),
    ])


def s4b():
    d = tsv(S2 / "shell_covariation.tsv")
    return pick(d, [
        ("context", "Comparison", "between: Trp subfamilies vs Phe subfamilies; otherwise the subfamily compared within "
                                  "(Trp- vs Phe-anchored members)"),
        ("pos", "Shell position", "Residue within 4.5 Å of the anchor side chain in >= 45% of AlphaFold models; "
                                  "relative to the Q-motif Gln (Q-n) or the anchor (UA+n)"),
        ("group_a", "Group A", "Phe-anchored group"),
        ("n_a", "Group A sequences", "One per genus and subfamily"),
        ("top_a", "Group A residues", "Most frequent residues (fraction)"),
        ("group_b", "Group B", "Trp-anchored group"),
        ("n_b", "Group B sequences", "One per genus and subfamily"),
        ("top_b", "Group B residues", "Most frequent residues (fraction)"),
        ("ref_residue", "Test residue", "Most frequent residue of group A"),
        ("frac_ref_a", "Test residue in A", "Fraction"),
        ("frac_ref_b", "Test residue in B", "Fraction"),
        ("p", "P", "Fisher's exact test, two-sided, test residue present vs absent"),
    ])


# --------------------------------------------------------------------- S5 ancestral states
def s5a():
    d = tsv(S3 / "asr_summary.tsv")
    return pick(d, [
        ("subfamily", "Subfamily", "Gene tree"),
        ("n_in", "Sequences", "Ingroup sequences"),
        ("tips", "Tip states", "Observed anchor residues (e.g. W340F60 = 340 Trp, 60 Phe)"),
        ("root_state", "Root state", "Maximum a posteriori state at the ingroup root (outgroup rooting)"),
        ("root_p", "Root posterior", "Marginal posterior probability of that state"),
        ("n_switches", "Switches", "Edges whose parent and child states differ"),
        ("to_W", "To Trp", "Switches to Trp"),
        ("from_W", "From Trp", "Switches from Trp"),
        ("anchor_rate", "Anchor site rate", "Empirical Bayes rate of the anchor column (IQ-TREE --rate)"),
        ("rate_percentile", "Rate percentile", "Percentile among all alignment columns (0 = slowest)"),
        ("frac_sites_slower_or_equal", "Columns slower or equal", "Fraction of columns with rate <= the anchor's"),
    ])


def s5b():
    d = tsv(S3 / "asr_switches.tsv")
    return pick(d, [
        ("subfamily", "Subfamily", "Gene tree"),
        ("from", "From", "Parent-node state"),
        ("to", "To", "Child state"),
        ("p_parent", "Parent posterior", "Marginal posterior of the parent state"),
        ("p_child", "Child posterior", "Marginal posterior of the child state (1 for tips)"),
        ("n_tips", "Clade size", "Sequences below the switch"),
        ("supergroups", "Supergroups", "Eukaryotic supergroups in the clade"),
        ("clade", "Clade", "Lowest taxon shared by the clade"),
        ("tip_states", "Tip states", "Observed anchor residues in the clade"),
    ])


# --------------------------------------------------------------------- S6 human variation
SITE = "UA: anchor; UA+1; Q-7/Q-6: Q-motif stacker; Q-4: Q-motif Pro; Q: Q-motif Gln; loop-n / loop+n: other anchor-loop residues (UA-8 to UA+3)"


def s6a():
    d = tsv(S4 / "variants_positions.tsv")
    return pick(d, [
        ("gene", "Gene", "HGNC symbol"),
        ("acc", "UniProt", "Accession"),
        ("anchor", "Anchor", "Anchor residue, UniProt numbering"),
        ("site", "Site", SITE),
        ("pos", "Position", "UniProt numbering"),
        ("source", "Source", "ClinVar or gnomAD v4"),
        ("hgvsp", "Protein change", "HGVS, canonical Ensembl transcript"),
        ("consequence", "Consequence", "VEP consequence"),
        ("clinvar", "ClinVar classification", "Clinical significance"),
        ("review", "ClinVar review status", "Review status"),
        ("gnomad_ac", "gnomAD allele count", "Allele count, exomes + genomes"),
        ("uniprot_match", "Matches UniProt", "Reference residue matches the UniProt sequence; only matching variants were analysed"),
        ("variant_id", "Variant", "GRCh38 chromosome-position-ref-alt"),
    ])


def s6b():
    d = tsv(S4 / "cancer_positions.tsv")
    return pick(d, [
        ("gene", "Gene", "HGNC symbol"),
        ("anchor", "Anchor", "Anchor residue, UniProt numbering"),
        ("site", "Site", SITE),
        ("pos", "Position", "UniProt numbering"),
        ("change", "Protein change", "As reported by cBioPortal"),
        ("type", "Mutation type", "cBioPortal classification"),
        ("study", "Study", "cBioPortal study (first seen)"),
        ("patient", "Patient", "Patient identifier; the same patient and change in several studies is counted once"),
        ("uniprot_match", "Matches UniProt", "Reference residue matches the UniProt sequence"),
        ("tumour_only", "Tumour", "FALSE = cell-line collection, excluded from counts"),
    ])


def s6c():
    g = tsv(S4 / "gnomad_missense_counts.tsv")
    c = tsv(S4 / "cancer_missense_counts.tsv").drop(columns="anchor")
    d = g.merge(c, on="gene", how="outer")
    return pick(d, [
        ("gene", "Gene", "HGNC symbol"),
        ("anchor", "Anchor", "Anchor residue, UniProt numbering"),
        ("transcript", "Transcript", "Canonical Ensembl transcript"),
        ("length", "Length (aa)", "Protein length"),
        ("missense_variants_total", "gnomAD missense variants", "Distinct missense variants in the protein (UniProt-matched)"),
        ("missense_sites_per_100aa", "gnomAD missense per 100 aa", "Distinct missense variants per 100 residues"),
        ("missense_at_UA", "gnomAD missense at anchor", "Distinct missense variants at the anchor"),
        ("missense_AC_at_UA", "gnomAD allele count at anchor", "Summed allele count"),
        ("missense_in_loop", "gnomAD missense in loop", "Distinct missense variants in the anchor loop"),
        ("numbering_mismatches", "Numbering mismatches", "gnomAD variants excluded because the reference residue did not match UniProt"),
        ("missense_patients", "Tumours with missense", "Tumours with a missense mutation in the protein (cell lines excluded)"),
        ("per_residue_mean", "Tumours per residue", "Mean over residues"),
        ("at_UA", "Tumours at anchor", "Tumours with a missense mutation at the anchor"),
        ("in_loop", "Tumours in loop", "Tumours with a missense mutation in the anchor loop"),
    ])


def s6d():
    d = tsv(S4 / "constraint.tsv")
    return pick(d, [
        ("dataset", "Dataset", "gnomAD v4 distinct missense variants, or cBioPortal tumours with a missense mutation"),
        ("site", "Site class", "UA, UA+1, Q-motif stacker, Q-motif Pro (Q-4), Q-motif Gln"),
        ("genes", "Genes", "Human DEAD-box proteins summed"),
        ("observed", "Observed", "Summed over genes"),
        ("expected", "Expected", "Per gene: total / protein length (average residue), summed"),
        ("ratio", "Observed/expected", "Ratio (Figure S1)"),
        ("p_depleted", "P (depleted)", "One-sided Poisson test"),
        ("p_enriched", "P (enriched)", "One-sided Poisson test"),
    ])


def s6e():
    d = tsv(S4 / "alphamissense_sites.tsv")
    return pick(d, [
        ("gene", "Gene", "HGNC symbol"),
        ("anchor", "Anchor", "Anchor residue, UniProt numbering"),
        ("site", "Site class", "UA, UA+1, Q-7/Q-6 (Q-motif stacker), Q-4 (Q-motif Pro), Q (Q-motif Gln)"),
        ("pos", "Position", "UniProt numbering"),
        ("residue", "Residue", "Wild-type residue"),
        ("rsasa", "Relative SASA", "In the AlphaFold model"),
        ("am_mean", "Mean AlphaMissense", "Mean pathogenicity over the 19 substitutions"),
        ("pct_all", "Percentile, all residues", "Among all residues of the protein"),
        ("pct_buried_aromatics", "Percentile, buried aromatics", "Among buried F/W/Y (relative SASA < 0.1) of the protein, anchor excluded"),
        ("n_buried_aromatics", "Buried aromatics", "Number in the control set"),
        ("buried_aromatics_median", "Buried aromatics median", "Median mean-AlphaMissense of the control set"),
    ])


def s6f():
    d = tsv(S4 / "alphamissense_swaps.tsv")
    cols = [("gene", "Gene", "HGNC symbol"), ("anchor", "Anchor", "Anchor residue, UniProt numbering")]
    cols += [(s, s.replace("->", "→"), f"AlphaMissense pathogenicity of {s.replace('->', '→')} at the anchor "
              "(likely benign < 0.34, likely pathogenic > 0.564)") for s in d.columns[2:]]
    return pick(d, cols)


# --------------------------------------------------------------------- workbook
def tables():
    c, e = census()
    n_seq, n_gen = len(e), e.genus.nunique()
    assert (n_seq, n_gen) == (36197, 2023), (n_seq, n_gen)   # numbers quoted in Results
    return [
        ("S1", "Nucleotide-bound SF2 helicase structures", [
            ("S1A", "Nucleotide sites", "All 535 adenine-nucleotide sites in 416 PDB entries of SF2 helicases, with the adenine "
             "stacker, the adenine-reading Gln, the motif I Lys and the upstream-anchor call.", *s1a()),
            ("S1B", "Family assignment", "Family label of each protein. Labels were set from UniProt names and Pfam keywords "
             "and every one was checked by hand.", *s1b()),
            ("S1C", "Anchor-N6 contact", "Geometry of the anchor-adenine N6 contact at all 119 DEAD-box sites with an inserted "
             "anchor (Figure 2A-B).", *s1c()),
        ]),
        ("S2", "DEAD-box proteins bound to other macromolecules", [
            ("S2A", "Partner complexes", "All 178 DEAD-box chains in 135 PDB entries that contain another polymer: anchor state "
             "and contacts with the anchor and its loop (deposited coordinates, cognate partner copy only).", *s2a()),
            ("S2B", "Free vs bound loop panel", "Eight DEAD-box proteins with free and partner-bound structures: per chain, the "
             "anchor-loop (UA-12 to UA+6) conformation relative to the nearest free chain, and partner contacts (Figure 2C-E).",
             *s2b()),
            ("S2C", "Anchors not inserted", "Chains of Table S2A in which the anchor is displaced or its state is ambiguous "
             "(Discussion).", *s2c()),
        ]),
        ("S3", "DEAD-box sequence census", [
            ("S3A", "Anchor residue by subfamily", f"Anchor residue in each eukaryotic subfamily (one sequence per genus and "
             f"subfamily; {n_seq:,} sequences, {n_gen:,} genera) and in all bacterial DEAD-box sequences (Figure 3A). "
             "UniProt release 2026_03, reference proteomes.", *s3a(c, e)),
            ("S3B", "Anchor spacing", "Distance from the anchor to the Q-motif Gln in eukaryotic sequences with an aromatic "
             "anchor (Figure 1B).", *s3b(e)),
            ("S3C", "Trp anchor by supergroup", "Percentage of genera with a Trp or Phe anchor in each subfamily and "
             "eukaryotic supergroup (Figure 4B).", *s3c()),
        ]),
        ("S4", "AlphaFold models", [
            ("S4A", "Anchor in AlphaFold models", "283 proteins selected (256 analysed): anchor contact with adenine transplanted "
             "from Vasa-AMP-PNP (PDB 2DB3) (Figure 3B-C).", *s4a()),
            ("S4B", "Anchor packing shell", "Residues packing around the anchor, compared between Trp- and Phe-anchored "
             "sequences (genus-balanced census).", *s4b()),
        ]),
        ("S5", "Ancestral anchor states", [
            ("S5A", "Gene trees", "Ancestral state reconstruction at the anchor column for twelve subfamilies (IQ-TREE 3, "
             "LG+G4) (Figure 4A).", *s5a()),
            ("S5B", "Anchor switches", "Every reconstructed change of the anchor residue, with the clade below it (Figure 4A).",
             *s5b()),
        ]),
        ("S6", "Human variation at the anchor and the Q motif", [
            ("S6A", "Germline variants", "ClinVar and gnomAD v4 variants at the anchor, UA+1, the Q-motif stacker, Pro and "
             "Gln, and the anchor loop of 37 human DEAD-box proteins (Figure 5A).", *s6a()),
            ("S6B", "Somatic mutations", "Mutations at the same positions in public cBioPortal studies (unique patient and "
             "change).", *s6b()),
            ("S6C", "Per-gene counts", "Missense variation per protein in gnomAD v4 and cBioPortal.", *s6c()),
            ("S6D", "Observed vs expected", "Missense variation at each site class compared with an average residue of the "
             "same protein (Figure 5C, Figure S1).", *s6d()),
            ("S6E", "AlphaMissense by site", "Mean AlphaMissense pathogenicity at each site class and its percentile "
             "(Figure 5B).", *s6e()),
            ("S6F", "AlphaMissense anchor swaps", "Predicted pathogenicity of substitutions at the anchor (Figure 5B).",
             *s6f()),
        ]),
    ]


HEAD = Font(name=FONT, bold=True, size=10)
BODY = Font(name=FONT, size=10)
TITLE = Font(name=FONT, bold=True, size=12)
FILL = PatternFill("solid", fgColor="E7EDF1")


def write_sheet(wb, sid, name, legend, df):
    ws = wb.create_sheet(sid)
    ws["A1"], ws["A1"].font = f"Table {sid}. {name}", TITLE
    ws["A2"], ws["A2"].font = legend + " Columns are defined in the README sheet.", BODY
    df = df.replace({np.nan: None})
    for j, h in enumerate(df.columns, 1):
        cell = ws.cell(row=4, column=j, value=h)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(df.itertuples(index=False), 5):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.bool_, bool)):
                v = bool(v)
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.floating):
                v = float(v)
            ws.cell(row=i, column=j, value=v).font = BODY
    for j, h in enumerate(df.columns, 1):
        vals = [len(str(x)) for x in df.iloc[:200, j - 1] if x is not None]
        ws.column_dimensions[get_column_letter(j)].width = min(45, max(9, len(h) * 0.55 + 2, *(vals or [0])) + 1)
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(df.shape[1])}{4 + len(df)}"


def main():
    groups = tables()
    wb = Workbook()
    rd = wb.active
    rd.title = "README"
    rd["A1"], rd["A1"].font = "Supplementary Tables S1-S6", TITLE
    rd["A2"] = ("An aromatic anchor, not the Q motif, marks the DEAD-box ATP site. One sheet per table; row 4 of each "
                "sheet holds the column headers. Residue positions are given relative to the Q-motif glutamine (Q-n) "
                "or the aromatic anchor (UA+n). Distances in Å. Per-sequence data (156,510 census sequences, gene "
                "trees, alignments) are in the data repository.")
    rd["A2"].font = BODY
    rd["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    rd.merge_cells("A2:D2")
    rd.row_dimensions[2].height = 60
    r = 4
    for gid, gtitle, subs in groups:
        rd.cell(row=r, column=1, value=f"Table {gid}. {gtitle}").font = TITLE
        r += 1
        for sid, name, legend, df, defs in subs:
            write_sheet(wb, sid, name, legend, df)
            c = rd.cell(row=r, column=1, value=sid)
            c.font, c.hyperlink = Font(name=FONT, bold=True, size=10, color="1F5F8B", underline="single"), f"#'{sid}'!A1"
            rd.cell(row=r, column=2, value=name).font = HEAD
            rd.cell(row=r, column=3, value=f"{len(df):,} rows").font = BODY
            rd.cell(row=r, column=4, value=legend).font = BODY
            rd.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
            for h, text in defs:
                rd.cell(row=r, column=2, value=h).font = BODY
                rd.cell(row=r, column=4, value=text).font = BODY
                rd.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
                r += 1
            r += 1
    for col, w in zip("ABCD", (8, 34, 11, 100)):
        rd.column_dimensions[col].width = w
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    for gid, _, subs in groups:
        print(gid, ", ".join(f"{sid} {len(df)}x{df.shape[1]}" for sid, _, _, df, _ in subs))
    print(OUT)


if __name__ == "__main__":
    main()
